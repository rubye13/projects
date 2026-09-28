# /// script
# requires-python = ">=3.12"
# dependencies = ["python-pptx==1.0.2", "pandas>=2.2"]
# ///
"""Собирает презентацию для Марии из цифр ноутбука экспериментов.

    uv run presentation/build_presentation.py

Цифры берутся из reports/experiments.csv и reports/feature_importance.csv,
поэтому после пересчёта ноутбука достаточно перезапустить скрипт.
Графики нарисованы фигурами, чтобы одинаково выглядеть в PowerPoint, Google Slides
и в превью на маке.
"""

import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
REPORTS = ROOT / "reports"
OUT = ROOT / "presentation" / "recsys_main_page.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
CORAL = RGBColor(0xF2, 0x5C, 0x54)
SLATE = RGBColor(0x5C, 0x6B, 0x8A)
LIGHT = RGBColor(0xEE, 0xF1, 0xF6)
MUTED = RGBColor(0xB8, 0xC2, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x28, 0x33)
PALE = RGBColor(0xC9, 0xD2, 0xE3)

TITLE_FONT = "Georgia"
BODY_FONT = "Arial"
# Заметки докладчика выключены: Quick Look на маке зависает на заметках от python-pptx,
# а превью пробелом - первое, что сделает человек со скачанным файлом
WITH_NOTES = "--notes" in sys.argv

FEATURE_LABELS = {
    "ui_days_since_last": "давно ли человек видел товар",
    "i_days_since_last": "давно ли с товаром что-то делали",
    "ui_carts": "клал ли товар в корзину",
    "ui_view_share": "доля товара в просмотрах человека",
    "i_views_7d": "просмотры товара за неделю",
    "u_days_since_first": "давно ли человек с нами",
    "ui_days_since_first": "когда впервые увидел товар",
    "i_carts": "корзины товара у всех",
    "uc_share": "интерес человека к категории",
    "i_available": "товар в наличии",
    "i_trend": "рост интереса к товару",
    "ui_in_last_session": "товар был в последнем визите",
    "i_cart_rate": "конверсия товара в корзину",
    "i_trans": "покупки товара",
    "u_cart_rate": "как часто человек кладёт в корзину",
    "i_views": "просмотры товара",
    "i_users": "посетители товара",
    "u_views": "просмотры человека",
    "i_price": "цена",
    "i_rank_in_cat": "место товара в категории",
    "ui_sessions": "в скольких визитах был товар",
    "ui_views": "сколько раз смотрел товар",
    "u_days_since_last": "давно ли человек заходил",
    "c_views": "просмотры категории",
}
MODEL_LABELS = {
    "популярное": "Популярное",
    "ALS": "ALS",
    "LightFM": "LightFM",
    "item-based": "Похожие товары",
    "корзина, потом свежесть": "Правило по истории",
    "CatBoost": "CatBoost",
}


# ---------- помощники ----------


def text(
    slide,
    left,
    top,
    width,
    height,
    runs,
    size=16,
    color=INK,
    bold=False,
    font=BODY_FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    spacing=None,
):
    """Текстовый блок. runs - строка или список абзацев, абзац - строка или список (текст, стиль)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        pieces = para if isinstance(para, list) else [(para, {})]
        for piece, style in pieces:
            run = p.add_run()
            run.text = piece
            f = run.font
            f.name = style.get("font", font)
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.italic = style.get("italic", False)
            f.color.rgb = style.get("color", color)
    return box


def card(slide, left, top, width, height, fill=LIGHT, radius=0.06):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def badge(slide, left, top, label, diameter=Inches(0.55), fill=CORAL, size=16):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.fill.background()
    circle.shadow.inherit = False
    frame = circle.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = WHITE
    return circle


def arrow(slide, x1, y1, x2, y2, color=SLATE, width=2.25):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    tail = line.line._get_or_add_ln()
    end = tail.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
        {"type": "triangle", "w": "med", "len": "med"},
    )
    tail.append(end)
    return line


def notes(slide, body):
    if WITH_NOTES:
        slide.notes_slide.notes_text_frame.text = body


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def title(slide, label, color=NAVY, top=Inches(0.55)):
    text(
        slide,
        Inches(0.7),
        top,
        Inches(11.9),
        Inches(0.9),
        label,
        size=36,
        color=color,
        bold=True,
        font=TITLE_FONT,
    )


def bar_chart(
    slide,
    left,
    top,
    width,
    height,
    categories,
    values,
    highlight,
    fmt,
    horizontal=False,
    base=NAVY,
    accent=CORAL,
    font_size=13,
    label_width=Inches(2.6),
):
    """Столбики из фигур с подсветкой главного.

    Нативный график PowerPoint, Google Slides и Quick Look рисуют по-разному: теряются
    подсветка, формат чисел и размер подписей. Прямоугольники и подписи везде одинаковые.
    """
    vmax = max(values)
    n = len(values)

    def bar(x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.shadow.inherit = False

    if horizontal:
        value_width = Inches(1.0)
        bar_space = width - label_width - value_width
        row = Emu(int(height / n))
        thick = Emu(int(row * 0.56))
        for i, (category, value) in enumerate(zip(categories, values)):
            y = top + Emu(int(row * i))
            color = accent if i in highlight else base
            length = Emu(max(int(bar_space * value / vmax), Inches(0.04)))
            bar(
                left + label_width,
                y + Emu(int((row - thick) / 2)),
                length,
                thick,
                color,
            )
            text(
                slide,
                left,
                y,
                label_width - Inches(0.15),
                row,
                category,
                size=font_size,
                align=PP_ALIGN.RIGHT,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            text(
                slide,
                left + label_width + length + Inches(0.1),
                y,
                value_width,
                row,
                fmt.format(value),
                size=font_size,
                bold=True,
                color=accent if i in highlight else INK,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        return

    label_height, value_height = Inches(0.45), Inches(0.4)
    plot_height = height - label_height - value_height
    col = Emu(int(width / n))
    bar_width = Emu(int(col * 0.5))
    base_y = top + value_height + plot_height
    for i, (category, value) in enumerate(zip(categories, values)):
        x = left + Emu(int(col * i))
        color = accent if i in highlight else base
        h = Emu(max(int(plot_height * value / vmax), Inches(0.03)))
        bar(x + Emu(int((col - bar_width) / 2)), base_y - h, bar_width, h, color)
        text(
            slide,
            x,
            base_y - h - value_height,
            col,
            value_height,
            fmt.format(value),
            size=font_size,
            bold=True,
            color=accent if i in highlight else INK,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        text(
            slide,
            x,
            base_y + Inches(0.1),
            col,
            label_height,
            category,
            size=font_size,
            align=PP_ALIGN.CENTER,
        )
    axis = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, left, base_y, left + width, base_y
    )
    axis.line.color.rgb = PALE
    axis.line.width = Pt(1)


# ---------- данные ----------


def load_numbers():
    table = pd.read_csv(REPORTS / "experiments.csv")
    test = table[table["выборка"] == "test"].set_index("модель")
    importance = pd.read_csv(REPORTS / "feature_importance.csv", index_col=0)[
        "важность"
    ]
    return test, importance


# ---------- слайды ----------


def slide_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    text(
        s,
        Inches(0.8),
        Inches(2.1),
        Inches(7.6),
        Inches(1.6),
        "Рекомендации на главной",
        size=48,
        color=WHITE,
        bold=True,
        font=TITLE_FONT,
    )
    text(
        s,
        Inches(0.8),
        Inches(3.95),
        Inches(7.2),
        Inches(1.0),
        "Что умеет модель, насколько она лучше популярного и как её подключить",
        size=20,
        color=PALE,
    )
    text(
        s,
        Inches(0.8),
        Inches(6.3),
        Inches(7),
        Inches(0.4),
        "Андрей Рубинштейн, команда рекомендаций",
        size=14,
        color=MUTED,
    )
    # три места на главной - мотив всей презентации
    for i, (label, fill) in enumerate(
        [
            ("1", CORAL),
            ("2", RGBColor(0xF5, 0x86, 0x7F)),
            ("3", RGBColor(0xF8, 0xAE, 0xA9)),
        ]
    ):
        left = Inches(8.9 + i * 1.35)
        tile = card(
            s, left, Inches(2.6), Inches(1.15), Inches(1.5), fill=fill, radius=0.12
        )
        tile.text_frame.text = ""
        text(
            s,
            left,
            Inches(2.95),
            Inches(1.15),
            Inches(0.8),
            label,
            size=40,
            color=WHITE,
            bold=True,
            font=TITLE_FONT,
            align=PP_ALIGN.CENTER,
        )
    text(
        s,
        Inches(8.9),
        Inches(4.3),
        Inches(3.85),
        Inches(0.5),
        "три места под товары",
        size=14,
        color=PALE,
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        (
            "Итоги проекта рекомендаций для главной страницы. Коротко: модель подбирает "
            "каждому посетителю три товара и в 2.5 раза точнее, чем показывать всем самое популярное."
        ),
    )


def slide_task(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Задача: рекомендовать каждому три товара")
    # макет главной
    page = card(
        s, Inches(0.7), Inches(1.75), Inches(5.2), Inches(4.2), fill=LIGHT, radius=0.04
    )
    page.text_frame.text = ""
    text(
        s,
        Inches(1.0),
        Inches(2.0),
        Inches(4.6),
        Inches(0.4),
        "Главная страница магазина",
        size=14,
        color=SLATE,
        bold=True,
    )
    for i in range(3):
        left = Inches(1.0 + i * 1.55)
        tile = card(
            s, left, Inches(2.65), Inches(1.4), Inches(1.9), fill=WHITE, radius=0.08
        )
        tile.text_frame.text = ""
        badge(s, left + Inches(0.42), Inches(3.0), str(i + 1))
        text(
            s,
            left,
            Inches(3.8),
            Inches(1.4),
            Inches(0.5),
            "товар для\nэтого человека",
            size=11,
            color=SLATE,
            align=PP_ALIGN.CENTER,
        )
    text(
        s,
        Inches(1.0),
        Inches(4.9),
        Inches(4.6),
        Inches(1.5),
        "Каждый посетитель видит свою тройку. Кто пришёл впервые, видит самые "
        "покупаемые товары",
        size=14,
        color=INK,
    )

    rows = [
        (
            "Бизнес-метрика",
            "Оборот от допродаж. Цель клиента +20%. Честно измерить её "
            "можно только A/B-тестом на живом трафике",
        ),
        (
            "Техническая метрика",
            "Precision@3: какая доля из трёх показанных товаров куплена "
            "за неделю после показа. По ней выбираем модель",
        ),
        (
            "Как проверяем",
            "Модель учится на прошлом и угадывает следующую неделю. "
            "Итоговая проверка на неделе 11-18 сентября, которую модель не видела",
        ),
    ]
    for i, (head, body) in enumerate(rows):
        top = Inches(1.75 + i * 1.65)
        badge(s, Inches(6.5), top + Inches(0.05), str(i + 1), fill=NAVY)
        text(
            s,
            Inches(7.25),
            top,
            Inches(5.4),
            Inches(0.4),
            head,
            size=18,
            bold=True,
            color=NAVY,
        )
        text(
            s, Inches(7.25), top + Inches(0.45), Inches(5.4), Inches(1.1), body, size=14
        )
    notes(
        s,
        (
            "Бизнес хочет оборот, но по истории продаж его не посчитать: нельзя знать, что человек "
            "купил бы без рекомендаций. Поэтому модель выбираем по Precision@3, а эффект в деньгах "
            "покажет A/B-тест."
        ),
    )


def slide_data(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Данные: три таблицы от клиента")
    tables = [
        (
            "events.csv",
            "2.76 млн строк",
            [
                ("timestamp", "число, время в мс"),
                ("visitorid", "число, id посетителя"),
                ("event", "текст: просмотр, корзина или покупка"),
                ("itemid", "число, id товара"),
                ("transactionid", "число, номер чека, только у покупок"),
            ],
        ),
        (
            "item_properties",
            "20.3 млн строк",
            [
                ("timestamp", "число, дата снимка (воскресенья)"),
                ("itemid", "число, id товара"),
                ("property", "текст, название свойства"),
                ("value", "текст, значение свойства"),
                (
                    "",
                    "Понятны только категория и наличие, остальное зашифровано. "
                    "Свойство 790 похоже на цену",
                ),
            ],
        ),
        (
            "category_tree",
            "1 669 строк",
            [
                ("categoryid", "число, id категории"),
                ("parentid", "число, родительская категория"),
                ("", "25 корневых разделов, до 5 уровней вложенности"),
            ],
        ),
    ]
    for i, (name, size, cols) in enumerate(tables):
        left = Inches(0.7 + i * 4.1)
        card(s, left, Inches(1.7), Inches(3.8), Inches(3.35)).text_frame.text = ""
        text(
            s,
            left + Inches(0.3),
            Inches(1.95),
            Inches(3.3),
            Inches(0.4),
            name,
            size=18,
            bold=True,
            color=NAVY,
            font=TITLE_FONT,
        )
        text(
            s,
            left + Inches(0.3),
            Inches(2.4),
            Inches(3.3),
            Inches(0.3),
            size,
            size=13,
            color=CORAL,
            bold=True,
        )
        paras = []
        for col, desc in cols:
            if col:
                paras.append([(col + "  ", {"bold": True, "color": NAVY}), (desc, {})])
            else:
                paras.append([(desc, {"italic": True, "color": SLATE})])
        text(
            s,
            left + Inches(0.3),
            Inches(2.9),
            Inches(3.3),
            Inches(2.6),
            paras,
            size=12,
            spacing=6,
        )
    facts = [
        ("4.5 месяца", "3 мая - 18 сентября 2015"),
        ("1.4 млн", "посетителей"),
        ("235 тыс.", "товаров в журнале"),
        ("22 тыс.", "покупок"),
    ]
    for i, (big, small) in enumerate(facts):
        left = Inches(0.7 + i * 3.075)
        text(
            s,
            left,
            Inches(5.55),
            Inches(2.9),
            Inches(0.6),
            big,
            size=28,
            bold=True,
            color=NAVY,
            font=TITLE_FONT,
        )
        text(
            s, left, Inches(6.15), Inches(2.9), Inches(0.4), small, size=13, color=SLATE
        )
    notes(
        s,
        (
            "Описания данных от клиента не было, смысл полей восстанавливали сами. Дубли в журнале "
            "убрали, свойства товаров берём на момент рекомендации, а не последние в файле."
        ),
    )


def slide_audience(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Покупателей мало, почти все новые")
    stats = [
        ("0.83%", "посетителей что-то покупают, в корзину кладут 2.7%"),
        ("71%", "посетителей сделали одно действие и ушли"),
        ("1 из 5", "покупателей недели заходил к нам за прошлые 8 недель"),
        ("0.61%", "покупок закрывал бейзлайн: один топ-3 для всех"),
    ]
    for i, (big, small) in enumerate(stats):
        left = Inches(0.7 + (i % 2) * 6.05)
        top = Inches(1.75 + (i // 2) * 2.25)
        card(s, left, top, Inches(5.75), Inches(1.95)).text_frame.text = ""
        text(
            s,
            left + Inches(0.4),
            top + Inches(0.3),
            Inches(2.4),
            Inches(1.2),
            big,
            size=44,
            bold=True,
            color=CORAL if i == 2 else NAVY,
            font=TITLE_FONT,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        text(
            s,
            left + Inches(2.9),
            top + Inches(0.35),
            Inches(2.6),
            Inches(1.3),
            small,
            size=15,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(
        s,
        Inches(0.7),
        Inches(6.35),
        Inches(11.9),
        Inches(0.7),
        "Отсюда устройство решения: персональная подборка для тех, у кого есть история, "
        "и самые покупаемые товары для всех остальных",
        size=16,
        color=NAVY,
        bold=True,
    )
    notes(
        s,
        (
            "Главная сложность данных - холодный старт. Большинство посетителей нам незнакомы, "
            "им модель показывает популярное. Персональные рекомендации получают самые ценные "
            "посетители: те, кто возвращается."
        ),
    )


def slide_recency(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Берут то, что недавно смотрели")
    text(
        s,
        Inches(0.7),
        Inches(1.5),
        Inches(7.6),
        Inches(0.6),
        "Сколько дней прошло с последнего просмотра товара и какая доля таких товаров "
        "окажется в корзине или покупке за неделю",
        size=13,
        color=SLATE,
    )
    bar_chart(
        s,
        Inches(0.7),
        Inches(2.15),
        Inches(7.6),
        Inches(4.05),
        ["до 1 дня", "1-3 дня", "3-7 дней", "1-2 недели", "2-4 недели", "4-8 недель"],
        [1.49, 0.88, 0.50, 0.37, 0.26, 0.14],
        highlight={0},
        fmt="{:.1f}%",
    )
    text(
        s,
        Inches(0.7),
        Inches(6.35),
        Inches(7.6),
        Inches(0.5),
        "В среднем по всем товарам из истории 0.4%",
        size=13,
        color=SLATE,
        bold=True,
    )

    card(s, Inches(8.9), Inches(1.9), Inches(3.75), Inches(2.05)).text_frame.text = ""
    text(
        s,
        Inches(9.2),
        Inches(2.1),
        Inches(3.2),
        Inches(0.8),
        "в 10 раз",
        size=36,
        bold=True,
        color=CORAL,
        font=TITLE_FONT,
    )
    text(
        s,
        Inches(9.2),
        Inches(2.9),
        Inches(3.2),
        Inches(1.0),
        "чаще берут товар, который смотрели вчера, чем месяц назад",
        size=14,
    )
    card(s, Inches(8.9), Inches(4.15), Inches(3.75), Inches(2.05)).text_frame.text = ""
    text(
        s,
        Inches(9.2),
        Inches(4.35),
        Inches(3.2),
        Inches(0.8),
        "в 8 раз",
        size=36,
        bold=True,
        color=NAVY,
        font=TITLE_FONT,
    )
    text(
        s,
        Inches(9.2),
        Inches(5.15),
        Inches(3.2),
        Inches(1.0),
        "чаще кладут в корзину или покупают то, что человек уже откладывал",
        size=14,
    )
    notes(
        s,
        (
            "Это главный сигнал для модели. От первого просмотра до покупки в медиане 10 минут, "
            "обычно всё происходит за один визит. Поэтому свежие просмотры и корзина важнее всего."
        ),
    )


def slide_features(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Какие факторы видит модель: 37 штук")
    groups = [
        (
            "9",
            "Человек и товар",
            [
                "сколько раз смотрел и клал в корзину",
                "как давно видел товар",
                "был ли товар в последнем визите",
                "интерес человека к категории товара",
            ],
        ),
        (
            "10",
            "Человек",
            [
                "сколько смотрит и покупает",
                "сколько было визитов и дней",
                "как давно заходил",
                "как часто кладёт в корзину",
            ],
        ),
        (
            "18",
            "Товар и категория",
            [
                "популярность за 8 недель и за последнюю неделю",
                "конверсия в корзину и в покупку",
                "цена, наличие, категория",
                "место товара в своей категории",
            ],
        ),
    ]
    for i, (count, head, items) in enumerate(groups):
        left = Inches(0.7 + i * 4.1)
        card(s, left, Inches(1.7), Inches(3.8), Inches(3.85)).text_frame.text = ""
        badge(
            s, left + Inches(0.3), Inches(1.95), count, diameter=Inches(0.75), size=18
        )
        text(
            s,
            left + Inches(1.25),
            Inches(2.1),
            Inches(2.4),
            Inches(0.5),
            head,
            size=18,
            bold=True,
            color=NAVY,
        )
        paras = [
            [("•  ", {"color": CORAL, "bold": True}), (item, {})] for item in items
        ]
        text(
            s,
            left + Inches(0.3),
            Inches(3.0),
            Inches(3.3),
            Inches(2.4),
            paras,
            size=14,
            spacing=8,
        )
    text(
        s,
        Inches(0.7),
        Inches(5.9),
        Inches(11.9),
        Inches(1.0),
        [
            [
                (
                    "Все факторы считаются только по прошлому: ",
                    {"bold": True, "color": NAVY},
                ),
                ("на дату рекомендации модель не знает ничего из будущего", {}),
            ],
            [
                (
                    "Генератор факторов Серёги тоже проверили: ",
                    {"bold": True, "color": NAVY},
                ),
                (
                    "время суток и топ-20 свойств товаров. Прироста не дали, в итоговую модель не вошли",
                    {},
                ),
            ],
        ],
        size=14,
        spacing=6,
    )
    notes(
        s,
        (
            "Факторы трёх уровней: про пару человек-товар, про человека и про товар. "
            "Самые полезные - про пару: как давно и как активно человек интересовался товаром."
        ),
    )


def slide_model(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Как модель выбирает три товара")
    steps = [
        ("История", "все просмотры, корзины и покупки человека за 8 недель"),
        ("Кандидаты", "до 50 последних товаров, с которыми он что-то делал"),
        (
            "CatBoost",
            "оценивает вероятность, что человек положит товар в корзину или купит",
        ),
        ("Три лучших", "товары с наибольшей вероятностью идут на главную"),
    ]
    width, gap, top = Inches(2.65), Inches(0.43), Inches(1.9)
    for i, (head, body) in enumerate(steps):
        left = Inches(0.7) + i * (width + gap)
        fill = CORAL if head == "CatBoost" else LIGHT
        card(s, left, top, width, Inches(2.15), fill=fill).text_frame.text = ""
        color = WHITE if head == "CatBoost" else NAVY
        text(
            s,
            left + Inches(0.25),
            top + Inches(0.25),
            width - Inches(0.5),
            Inches(0.5),
            head,
            size=18,
            bold=True,
            color=color,
        )
        text(
            s,
            left + Inches(0.25),
            top + Inches(0.8),
            width - Inches(0.5),
            Inches(1.3),
            body,
            size=13,
            color=WHITE if head == "CatBoost" else INK,
        )
        if i < len(steps) - 1:
            arrow(
                s,
                left + width + Inches(0.06),
                top + Inches(1.07),
                left + width + gap - Inches(0.06),
                top + Inches(1.07),
            )

    card(
        s, Inches(0.7), Inches(4.45), Inches(11.9), Inches(1.0), fill=WHITE
    ).line.color.rgb = PALE
    text(
        s,
        Inches(1.0),
        Inches(4.55),
        Inches(11.3),
        Inches(0.8),
        [
            [
                (
                    "Нет истории или меньше трёх своих товаров?  ",
                    {"bold": True, "color": NAVY},
                ),
                ("Свободные места занимают самые покупаемые товары за 8 недель", {}),
            ],
        ],
        size=14,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        Inches(0.7),
        Inches(5.9),
        Inches(11.9),
        Inches(1.0),
        [
            [
                ("CatBoost ", {"bold": True, "color": NAVY}),
                (
                    "- это сотни небольших деревьев решений, каждое следующее исправляет ошибки "
                    "предыдущих. Учится на прошлых неделях: для каждого товара из истории человека "
                    "известно, положил ли он его потом в корзину или купил",
                    {},
                ),
            ],
        ],
        size=14,
    )
    notes(
        s,
        (
            "Модель не придумывает товары с нуля, а ранжирует то, чем человек уже интересовался. "
            "Мы проверили и модели, которые предлагают похожие товары: на трёх местах они проигрывают."
        ),
    )


def slide_importance(prs, importance):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Что влияет на решение модели")
    top = importance.head(8)
    bar_chart(
        s,
        Inches(0.7),
        Inches(1.75),
        Inches(7.9),
        Inches(4.7),
        [FEATURE_LABELS.get(f, f) for f in top.index],
        list(top.values),
        highlight={0, 1, 2},
        fmt="{:.1f}%",
        horizontal=True,
        label_width=Inches(3.7),
    )
    text(
        s,
        Inches(0.7),
        Inches(6.65),
        Inches(7.9),
        Inches(0.4),
        "Вклад фактора в предсказания модели, %",
        size=12,
        color=SLATE,
    )

    points = [
        ("Наличие", "товар, которого нет на складе, модель наверх не ставит"),
        ("Свежесть", "чем свежее человек смотрел товар, тем выше он в подборке"),
        ("Корзина", "товар, который человек уже откладывал, поднимается наверх"),
    ]
    for i, (head, body) in enumerate(points):
        y = Inches(1.8 + i * 1.6)
        badge(s, Inches(9.0), y, str(i + 1))
        text(
            s,
            Inches(9.75),
            y,
            Inches(2.9),
            Inches(0.4),
            head,
            size=18,
            bold=True,
            color=NAVY,
        )
        text(s, Inches(9.75), y + Inches(0.45), Inches(2.9), Inches(1.0), body, size=14)
    notes(
        s,
        (
            "Коротко: рекомендуем то, что есть в наличии, что человек недавно смотрел и откладывал, "
            "и что хорошо кладут в корзину другие. Номер категории и цена почти не влияют."
        ),
    )


def slide_result(prs, test):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Итог на тестовой неделе")
    best = test.loc["CatBoost"]
    base = test.loc["популярное"]
    ratio = best["P@3, %"] / base["P@3, %"]
    stats = [
        (f"{best['P@3, %']:.1f}%", "Precision@3 модели", CORAL),
        (f"×{ratio:.1f}", f"к популярному ({base['P@3, %']:.2f}%)", NAVY),
        (
            f"{best['покрытие покупок, %']:.1f}%",
            "покупок недели попали в подборки, " "у бейзлайна было 0.61%",
            NAVY,
        ),
    ]
    for i, (big, small, color) in enumerate(stats):
        top = Inches(1.7 + i * 1.65)
        card(s, Inches(0.7), top, Inches(4.6), Inches(1.45)).text_frame.text = ""
        text(
            s,
            Inches(1.0),
            top + Inches(0.12),
            Inches(2.0),
            Inches(1.2),
            big,
            size=36,
            bold=True,
            color=color,
            font=TITLE_FONT,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        text(
            s,
            Inches(3.0),
            top + Inches(0.15),
            Inches(2.15),
            Inches(1.15),
            small,
            size=13,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    order = test.sort_values("P@3, %", ascending=False, kind="stable")
    bar_chart(
        s,
        Inches(5.8),
        Inches(2.1),
        Inches(6.9),
        Inches(4.1),
        [MODEL_LABELS.get(m, m) for m in order.index],
        list(order["P@3, %"]),
        highlight={list(order.index).index("CatBoost")},
        fmt="{:.2f}%",
        horizontal=True,
        base=MUTED,
        label_width=Inches(2.3),
    )
    text(
        s,
        Inches(5.8),
        Inches(1.65),
        Inches(6.9),
        Inches(0.4),
        "Precision@3 разных моделей на тестовой неделе 11-18 сентября",
        size=13,
        color=SLATE,
    )
    text(
        s,
        Inches(5.8),
        Inches(6.4),
        Inches(6.9),
        Inches(0.6),
        "Потолок для этих данных около 5.4%: остальные покупки - товары, "
        "которых человек раньше не видел",
        size=12,
        color=SLATE,
    )
    notes(
        s,
        (
            "Цифры маленькие в абсолюте, потому что большинство покупок - товары, которые человек "
            "увидел впервые в тот же визит. Из того, что вообще можно угадать заранее, модель "
            "угадывает большую часть."
        ),
    )


def slide_next(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    title(s, "Что готово и что дальше", color=WHITE)
    done = [
        "Исследование данных",
        "Факторы для модели",
        "Эксперименты: 6 подходов",
        "Сервис рекомендаций (MVP)",
        "Docker-образ и документация",
    ]
    text(
        s,
        Inches(0.8),
        Inches(1.7),
        Inches(5.3),
        Inches(0.5),
        "Готово",
        size=20,
        bold=True,
        color=CORAL,
    )
    for i, item in enumerate(done):
        y = Inches(2.35 + i * 0.72)
        badge(s, Inches(0.8), y, "✓", diameter=Inches(0.45), size=14)
        text(
            s,
            Inches(1.45),
            y + Inches(0.05),
            Inches(4.6),
            Inches(0.45),
            item,
            size=16,
            color=WHITE,
        )

    text(
        s,
        Inches(6.9),
        Inches(1.7),
        Inches(5.7),
        Inches(0.5),
        "Дальше",
        size=20,
        bold=True,
        color=CORAL,
    )
    nxt = [
        (
            "A/B-тест на части трафика",
            "измерить прирост оборота, ради которого всё затевалось",
        ),
        (
            "Пересчёт внутри визита",
            "обновлять тройку после первых просмотров, так ловим "
            "покупки того же визита",
        ),
        ("Похожие товары", "добавить в кандидаты товары, похожие на просмотренные"),
    ]
    for i, (head, body) in enumerate(nxt):
        y = Inches(2.35 + i * 1.2)
        text(
            s,
            Inches(6.9),
            y,
            Inches(5.7),
            Inches(0.4),
            head,
            size=16,
            bold=True,
            color=WHITE,
        )
        text(
            s,
            Inches(6.9),
            y + Inches(0.4),
            Inches(5.7),
            Inches(0.7),
            body,
            size=13,
            color=PALE,
        )
    text(
        s,
        Inches(0.8),
        Inches(6.35),
        Inches(11.8),
        Inches(0.5),
        "Сервис отвечает за 17 мс, держит около 400 запросов в секунду и не падает "
        "на некорректных запросах",
        size=14,
        color=MUTED,
    )
    notes(
        s,
        (
            "Все этапы плана закрыты. Следующий шаг с клиентом - A/B-тест: только он покажет, "
            "даёт ли модель обещанные +20% к обороту от допродаж."
        ),
    )


def build(only=None):
    test, importance = load_numbers()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    props = prs.core_properties
    props.author = props.last_modified_by = "Андрей Рубинштейн"
    props.title = "Рекомендации на главной"
    makers = [
        lambda: slide_title(prs),
        lambda: slide_task(prs),
        lambda: slide_data(prs),
        lambda: slide_audience(prs),
        lambda: slide_recency(prs),
        lambda: slide_features(prs),
        lambda: slide_model(prs),
        lambda: slide_importance(prs, importance),
        lambda: slide_result(prs, test),
        lambda: slide_next(prs),
    ]
    for i, make in enumerate(makers, start=1):
        if only is None or i == only:
            make()
    return prs


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--each":
        # по файлу на слайд, чтобы посмотреть каждый через Quick Look
        out_dir = Path(sys.argv[2])
        out_dir.mkdir(parents=True, exist_ok=True)
        for n in range(1, 11):
            build(only=n).save(out_dir / f"slide_{n:02d}.pptx")
        print("слайды по отдельности:", out_dir)
    else:
        build().save(OUT)
        print("сохранено:", OUT)
